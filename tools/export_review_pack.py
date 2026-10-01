#!/usr/bin/env python3
"""Build the clinical review pack: every distractor card from content_drafts/.

One row per drafted card, with its source and the number of cases it reached,
grouped by specialty, so a reviewer can work in a spreadsheet instead of
reading JSON. Reviewer columns (veredicto, correccion, fuente verificada) are
left blank, with a dropdown for the verdict.

Outputs in <out_dir>:
  revision_clinica.xlsx   sheets: LEEME, Prioridad (riesgo letal), one per
                          specialty, Correcciones, Resumen
  revision_clinica.csv    same rows, flat

Usage: python3 tools/export_review_pack.py <cases_dir> <drafts_dir> <out_dir>
Requires openpyxl.
"""
import csv
import json
import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SPECIALTY = {
    "CARD": "Cardiología", "INT": "Medicina Interna", "GYN": "Ginecología",
    "OBS": "Obstetricia", "OBG": "Obstetricia", "PED": "Pediatría",
    "SURG": "Cirugía y Trauma", "GAST": "Gastroenterología",
    "NEUR": "Neurología", "PSYC": "Psiquiatría", "PSYCH": "Psiquiatría",
    "INF": "Infectología", "ENDO": "Endocrinología", "OPH": "Oftalmología",
    "DERM": "Dermatología", "PREV": "Medicina Preventiva",
    "IM": "Medicina Interna", "ORL": "Otorrinolaringología",
    "STATS": "Bioestadística", "ENGL": "Inglés técnico",
}
HEADERS = ["#", "Lote", "Especialidad", "Tema", "Casos afectados", "Card ID",
           "Riesgo letal", "Categoría", "Texto de la carta (se descarta)",
           "Comentario del mentor", "Fuente citada",
           "Veredicto", "Corrección sugerida", "Fuente verificada"]
WIDTHS = [5, 8, 18, 30, 9, 10, 9, 18, 60, 60, 45, 13, 45, 11]
# Cards moved from keep to discard in this PR (family, card_id).
CORRECTED = [
    ("PROC_SURG_APPENDICITIS_001", "c_3"),
    ("PROC_PED_RESPIRATORY_CRUP_001", "c_2"),
    ("PROC_SURG_APPENDICITIS_ALVARADO_001", "c_8"),
]


def split_family(family):
    code = family.split("_")[1] if family.startswith("PROC_") else family.split("_")[0]
    topic = re.sub(r"^PROC_[A-Z]+_|_\d+$", "", family).replace("_", " ").title()
    return SPECIALTY.get(code, "Otras"), topic


def count_cases(cases_dir, family):
    return len(list(Path(cases_dir).glob(f"CASE_{family}_*.json")))


def load_rows(cases_dir, drafts_dir):
    rows = []
    for path in sorted(Path(drafts_dir).glob("distractores_lote*.json")):
        lote = re.search(r"lote(\w+)\.json", path.name).group(1)
        for family, cards in json.loads(path.read_text(encoding="utf-8"))["familias"].items():
            spec, topic = split_family(family)
            n = count_cases(cases_dir, family)
            for c in cards:
                if c.get("reserva"):
                    continue
                rows.append([
                    0, lote, spec, topic, n, c["card_id"],
                    "Sí" if c.get("lethal_risk") else "", c["category"],
                    c["card_text"], c["comment"], c.get("fuente", ""),
                    "", "", "",
                ])
    rows.sort(key=lambda r: (r[2], r[3], r[1], r[5]))
    for i, r in enumerate(rows, 1):
        r[0] = i
    return rows


def load_corrections(cases_dir):
    out = []
    for family, card_id in CORRECTED:
        path = sorted(Path(cases_dir).glob(f"CASE_{family}_*.json"))[0]
        case = json.loads(path.read_text(encoding="utf-8"))
        # Alvarado's adenitis card id differs; match by text when not found.
        card = next((c for c in case["card_stream"] if c["card_id"] == card_id), None)
        if card is None or card["expected_action"] != "discard":
            card = next(c for c in case["card_stream"]
                        if c["card_text"].startswith(("Adenitis", "Epiglotitis", "Embarazo")))
        spec, topic = split_family(family)
        out.append([spec, topic, card["card_id"], card["card_text"],
                    card["scoring"]["vazquez_comment"],
                    "Estaba como 'aceptar'; describe un diagnóstico que no corresponde al paciente. Ahora es 'descartar'.",
                    "", "", ""])
    return out


def style_sheet(ws, header_row=1):
    head = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def add_review_sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    style_sheet(ws)
    dv = DataValidation(type="list", formula1='"Aprobar,Corregir,Rechazar"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    last = max(ws.max_row, 2)
    dv.add(f"L2:L{last}")
    dv2.add(f"N2:N{last}")
    ws.auto_filter.ref = f"A1:N{last}"
    return ws


LEEME = [
    "Revisión clínica de cartas nuevas de Dr. Swipe",
    "",
    "Qué es: las cartas de DESCARTE (distractores) agregadas en los lotes 1 a 6. Cada una es una afirmación falsa o inadecuada que el jugador debe deslizar a la izquierda.",
    "Qué se pide: confirmar para cada carta que (1) la afirmación es de verdad incorrecta o inadecuada para ese paciente y (2) el comentario del mentor es clínicamente correcto, tanto si el jugador acierta como si falla.",
    "",
    "Cómo usar el libro:",
    "  1. 'Prioridad (riesgo letal)': empieza aquí. Son las cartas cuyo error sería peligroso; un fallo aquí pesa más.",
    "  2. Una hoja por especialidad: revisa las demás cartas por área.",
    "  3. Columna Veredicto: Aprobar / Corregir / Rechazar. Si corriges, escribe el texto correcto en 'Corrección sugerida'.",
    "  4. 'Fuente verificada': marca Sí solo si abriste la guía citada y la afirmación está ahí.",
    "  5. 'Correcciones': 3 diagnósticos diferenciales que estaban como 'aceptar' y ahora son 'descartar' (afectan 4 casos).",
    "",
    "Límite importante: las fuentes citadas se redactaron de memoria; los textos originales no se consultaron. Ninguna cita está verificada hasta que marques la columna N.",
    "'Casos afectados' indica a cuántos casos del juego llega cada carta, para pesar el impacto.",
    "Pendientes de decisión clínica aparte: roséola (15 casos) y el umbral de SatO2 en la GPC mexicana.",
]


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        return 2
    cases_dir, drafts_dir, out_dir = sys.argv[1:4]
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    rows = load_rows(cases_dir, drafts_dir)
    corrections = load_corrections(cases_dir)

    with open(Path(out_dir) / "revision_clinica.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "LEEME"
    for line in LEEME:
        ws.append([line])
    ws.column_dimensions["A"].width = 140
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    add_review_sheet(wb, "Prioridad (riesgo letal)", [r for r in rows if r[6] == "Sí"])
    by_spec = OrderedDict()
    for r in rows:
        by_spec.setdefault(r[2], []).append(r)
    for spec, group in by_spec.items():
        add_review_sheet(wb, spec, group)

    wc = wb.create_sheet("Correcciones")
    wc.append(["Especialidad", "Tema", "Card ID", "Texto de la carta",
               "Comentario nuevo", "Qué cambió", "Veredicto",
               "Corrección sugerida", "Fuente verificada"])
    for r in corrections:
        wc.append(r)
    for col, w in zip("ABCDEFGHI", [18, 30, 10, 60, 60, 45, 13, 45, 11]):
        wc.column_dimensions[col].width = w
    style_sheet(wc)

    wr = wb.create_sheet("Resumen")
    wr.append(["Especialidad", "Cartas", "Con riesgo letal"])
    total = lethal = 0
    for spec, group in by_spec.items():
        n, l = len(group), sum(1 for r in group if r[6] == "Sí")
        wr.append([spec, n, l])
        total += n
        lethal += l
    wr.append(["Total", total, lethal])
    wr.column_dimensions["A"].width = 28
    style_sheet(wr)
    wr.cell(row=wr.max_row, column=1).font = Font(bold=True)

    wb.save(Path(out_dir) / "revision_clinica.xlsx")
    c = Counter(r[1] for r in rows)
    print(f"{total} cartas ({lethal} con riesgo letal) en {len(by_spec)} especialidades; lotes: {dict(sorted(c.items()))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
