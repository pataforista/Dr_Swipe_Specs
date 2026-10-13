#!/usr/bin/env python3
"""Write the 2026-10-13 verdicts into REVISION_CLINICA (csv and xlsx).

Run after tools/export_review_pack.py, which rebuilds the pack with empty
verdict columns. Rows are matched by the start of the card text.
"Fuente verificada": Sí = the statement was read in PubMed/Consensus;
Parcial = the source was located but its text was not readable for this
statement; No = no source located (Insufficient data).

Usage: python3 tools/fill_review_verdicts.py
"""
import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PACK = ROOT / "REVISION_CLINICA"

A, C = "Aprobar", "Corregir"
NOTE_APPLIED = "Aplicada el 2026-10-13 (tools/apply_clinical_review_2026_10.py)"
V = [
    ("Fibrinólisis sistémica como primera opción en este choque", C, "Sí", f"Comentario: ICP inmediata; fibrinólisis solo si la ICP no llega a tiempo. {NOTE_APPLIED}"),
    ("Bolos repetidos de cristaloides de 1 litro", A, "Parcial", ""),
    ("Bloqueador beta intravenoso para reducir la frecuencia", A, "Parcial", ""),
    ("Tiempo puerta-balón de 4 horas", C, "Sí", f"Comentario: umbral de 120 min desde el diagnóstico, no 90 min desde el primer contacto. {NOTE_APPLIED}"),
    ("Oxígeno en mascarilla con reservorio", A, "Parcial", "ATLS 10a ed. es un libro, no indexado en PubMed"),
    ("Bolos de cristaloides hasta normalizar la presión", A, "Parcial", "ATLS 10a ed. es un libro, no indexado en PubMed"),
    ("Colecistectomía urgente como tratamiento inicial", A, "Parcial", "TG18 (doi:10.1002/jhbp.509): drenaje biliar en colangitis moderada/grave"),
    ("Antibiótico hasta después de la CPRE", A, "Sí", "TG18: tratamiento inmediato sin esperar diagnóstico definitivo; SSC 2021: antimicrobiano en 1 h"),
    ("Drenaje biliar diferido 72 horas", A, "Parcial", "TG18 (doi:10.1002/jhbp.509)"),
    ("Metoclopramida para acelerar el tránsito", C, "Sí", f"Solo la fuente (la citada no trata metoclopramida). {NOTE_APPLIED}"),
    ("Cristaloides en bolos sin límite", A, "Parcial", "ATLS 10a ed. no indexado"),
    ("Intubación con ventilación a presión positiva antes", A, "Parcial", "ATLS 10a ed. no indexado"),
    ("Trasladar a tomografía antes de estabilizar", A, "Parcial", "ATLS 10a ed. no indexado"),
    ("Suspender la insulina al bajar la glucosa", A, "Sí", "ADA 2024 (doi:10.2337/dci24-0032) localizada; texto confirmado en JBDS: insulina fija hasta resolver la cetosis y glucosa al bajar"),
    ("Enema con bario en la fase aguda", A, "Parcial", "ASCRS 2020 (doi:10.1097/DCR.0000000000001679) sin resumen"),
    ("Evitar la heparina por el riesgo de sangrado intestinal", A, "Sí", "WSES 2022 (doi:10.1186/s13017-022-00443-x): heparina temprana mejora la supervivencia"),
    ("Vasopresores en dosis altas para sostener la presión", A, "Sí", "WSES 2022: vasopresores con cautela; eliminarlos es medida primaria"),
    ("Observación y ayuno por 24 horas", A, "Parcial", "WSES 2022"),
    ("Transfusión hasta hemoglobina mayor de 10", C, "Sí", f"Meta 7-8 g/dL (Baveno VII 6.2), no 7-9. {NOTE_APPLIED}"),
    ("Propranolol en dosis plena durante la hemorragia activa", A, "Sí", "Baveno VII 7.7: reducir o suspender BBNS con PAS <90 o PAM <65"),
    ("Esperar la tomografía y el resultado del LCR", A, "Parcial", "IDSA 2004 (doi:10.1086/425368) y ESCMID 2016 (doi:10.1016/j.cmi.2016.01.007) localizadas sin resumen"),
    ("Esperar los Cultivos", A, "Sí", "SSC 2021 (doi:10.1007/s00134-021-06506-y): antimicrobiano en 1 h; cultivos sin retrasarlo"),
    ("Doxiciclina en el Embarazo", A, "Parcial", "CDC 2021 (doi:10.15585/mmwr.rr7004a1) localizada"),
    ("Nitroglicerina sublingual repetida hasta aliviar el dolor", A, "Parcial", ""),
    ("Esperar la elevación de la troponina", A, "Parcial", ""),
    ("Bloqueador beta intravenoso en bolo desde el ingreso", A, "Parcial", ""),
    ("Salmeterol en Monoterapia", A, "Sí", "SMART (doi:10.1378/chest.129.1.15): más muertes por asma con salmeterol"),
    ("Propranolol: Útil para controlar la taquicardia", A, "Parcial", ""),
    ("Nifedipino sublingual de acción corta", A, "Parcial", ""),
    ("Trombólisis en TEP de Bajo Riesgo", A, "Parcial", ""),
    ("Observar al murciélago por 10 días", A, "Parcial", "ACIP 2008 (PMID 18496505) localizada; MMWR 2026: murciélago es la principal fuente de rabia humana en EE. UU."),
    ("Esperar el resultado de laboratorio del animal", A, "Parcial", "ACIP 2008 localizada"),
    ("Nimodipino intravenoso en bolo", A, "Parcial", "AHA/ASA 2023 (doi:10.1161/STR.0000000000000436) localizada; vía oral no confirmada en resumen"),
    ("Aspirina junto con el rTPA", C, "Sí", f"Solo la fuente: se añade ARTIS (doi:10.1016/S0140-6736(12)60949-0). {NOTE_APPLIED}"),
    ("Misoprostol como único tratamiento del aborto séptico", A, "Parcial", "ACOG PB 200 (doi:10.1097/AOG.0000000000002899) sin mención de infección en el resumen"),
    ("Tocólisis de Rutina: Indicada en todo DPPNI", C, "Sí", f"Solo la fuente (Medscape no es fuente primaria). {NOTE_APPLIED}"),
    ("Metotrexato en la Paciente Inestable", C, "Sí", f"El hemoperitoneo por sí solo no es contraindicación absoluta; sí la inestabilidad o la rotura (ACOG PB 193). {NOTE_APPLIED}"),
    ("Ergonovina en Preeclampsia", A, "Sí", "FIGO 2022 (doi:10.1002/ijgo.14116): oxitocina primera línea; ergot tras excluir hipertensión"),
    ("Umbral de Actuación", A, "Sí", "FIGO 2022: definición clásica >500 mL vaginal, >1000 mL cesárea"),
    ("Terbutalina: Útil para controlar la atonía", A, "No", "Insufficient data: sin fuente específica localizada; el mecanismo (tocolítico) es farmacología básica"),
    ("Tocólisis: Inhibir las contracciones", A, "No", "Insufficient data: fuente no localizada en PubMed"),
    ("Diazepam como anticonvulsivo de primera línea", A, "No", "ACOG PB 222 no localizada en PubMed"),
    ("Vacuna SRP a la embarazada", A, "Sí", "ACIP 2013 (PMID 23760231): evitar el embarazo tras vacuna con rubéola"),
    ("Choque por Deshidratación", A, "No", "Plan C OMS/AIEPI no indexado"),
    ("Esperar el Apgar del minuto 1", A, "Parcial", "AHA 2020 (doi:10.1161/CIR.0000000000000902) localizada"),
    ("Adrenalina intravenosa a 0.1 mg/kg", A, "Sí", "IV 0.01-0.03 mg/kg (Ramsie 2024, doi:10.1159/000535502; ILCOR)"),
    ("Iniciar compresiones torácicas con frecuencia", A, "Sí", "FC <60 tras 30 s de ventilación (ILCOR; Ramachandran 2022)"),
    ("Cirugía urgente antes de corregir", A, "No", "La fuente citada ('AAP Clinical Report') no se localizó; Insufficient data"),
    ("Receta de antidepresivo por 3 meses", A, "No", "APA 2010 no localizada en PubMed"),
    ("Escala SAD PERSONS como único criterio", C, "Sí", f"Solo la fuente: Steeg 2018 (AUC 51%). {NOTE_APPLIED}"),
    ("Antidepresivo en monoterapia durante el episodio maníaco", A, "Parcial", "CANMAT/ISBD localizada (doi:10.1111/bdi.13135)"),
    ("Reiniciar el antipsicótico a las 24 horas", C, "Parcial", f"Periodo sin fármaco y cambio de agente (Pileggi 2016); '2 semanas' queda como 'suele recomendarse'. {NOTE_APPLIED}"),
    ("Recuento de neutrófilos innecesario", A, "No", "APA 2020 no localizada en PubMed"),
]
CORR = [
    ("Embarazo Ectópico Roto", C, "", f"El comentario ahora menciona que no hay sangrado transvaginal. Pendiente para el revisor: en mujer en edad fértil, ¿añadir prueba de embarazo? (sin fuente verificada). {NOTE_APPLIED}"),
    ("Epiglotitis Aguda", C, "", f"El escenario de CRUP_001_002 no menciona tos perruna; el comentario ya no la afirma. {NOTE_APPLIED}"),
    ("Adenitis Mesentérica", A, "", "José, 32 años: el comentario corresponde al paciente"),
]


def lookup(text, table):
    for prefix, verdict, src, note in table:
        if text.startswith(prefix):
            return verdict, note, src
    return None


def main() -> int:
    path = PACK / "revision_clinica.csv"
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig")))
    hit = 0
    for r in rows:
        if r["Riesgo letal"] != "Sí":
            continue
        v = lookup(r["Texto de la carta (se descarta)"], V)
        if v:
            r["Veredicto"], r["Corrección sugerida"], r["Fuente verificada"] = v
            hit += 1
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print(f"csv: {hit} filas con veredicto")

    wb = openpyxl.load_workbook(PACK / "revision_clinica.xlsx")
    filled = 0
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]]
        if "Veredicto" not in header:
            continue
        col = {h: i for i, h in enumerate(header)}
        text_col = col.get("Texto de la carta (se descarta)", col.get("Texto de la carta"))
        table = CORR if ws.title == "Correcciones" else V
        for row in ws.iter_rows(min_row=2):
            text = row[text_col].value or ""
            if ws.title != "Correcciones" and "Riesgo letal" in col and row[col["Riesgo letal"]].value != "Sí":
                continue
            v = lookup(text, table)
            if v:
                row[col["Veredicto"]].value, row[col["Corrección sugerida"]].value, row[col["Fuente verificada"]].value = v
                filled += 1
    wb.save(PACK / "revision_clinica.xlsx")
    print(f"xlsx: {filled} celdas de veredicto llenas")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
